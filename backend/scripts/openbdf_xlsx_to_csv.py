"""Export the data sheets of an OpenBDF workbook to CSV, one per sheet, by field code.

The API reads these CSVs (OPENBDF_DATA_DIR) instead of the workbook: the lmc-openbdf
parser does not read v2 working files yet, and openpyxl stays out of the runtime image.

    uv run --no-project --with openpyxl \
        python scripts/openbdf_xlsx_to_csv.py openbdf_at.xlsx data/openbdf
"""

import csv
import sys
from pathlib import Path

import openpyxl

# Sheets laid out one row per record: header rows labelled in column A, then data rows.
ROW_SHEETS = {
    "Bill of Materials": "bill_of_materials.csv",
    "LCA Result (Wide)": "lca_results_wide.csv",
    "Project material database": "project_material_database.csv",
}
# General Information has one row per field: code in G (or H if updated), value in Q.
GI_CODE, GI_UPDATED_CODE, GI_VALUE = 6, 7, 16


def cell_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def export_row_sheet(sheet, out: Path) -> int:
    headers: dict[str, tuple] = {}
    rows: list[tuple] = []
    for row in sheet.iter_rows(values_only=True):
        if not rows and row[0] is not None:
            headers[cell_text(row[0])] = row
            continue
        if cell_text(row[1]) in ("", "…"):
            continue
        rows.append(row)

    field, updated = headers["Field Code"], headers.get("UPDATED Field Code")
    codes = [
        cell_text(updated[i]) if updated and updated[i] else cell_text(field[i])
        for i in range(len(field))
    ]
    keep = [i for i, code in enumerate(codes) if i > 0 and code]
    duplicates = {
        codes[i] for i in keep if [codes[j] for j in keep].count(codes[i]) > 1
    }
    if duplicates:
        raise ValueError(f"{sheet.title}: duplicate field codes {sorted(duplicates)}")

    with out.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(codes[i] for i in keep)
        writer.writerows([cell_text(row[i]) for i in keep] for row in rows)
    return len(rows)


def export_general_information(sheet, out: Path) -> int:
    with out.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["field_code", "value"])
        count = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            code = cell_text(row[GI_UPDATED_CODE]) or cell_text(row[GI_CODE])
            if code:
                writer.writerow([code, cell_text(row[GI_VALUE])])
                count += 1
    return count


def main(workbook: Path, out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    # data_only: the cached values, which is what the contributor's tool wrote.
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    for sheet_name, file_name in ROW_SHEETS.items():
        count = export_row_sheet(wb[sheet_name], out_dir / file_name)
        print(f"{file_name}: {count} rows")
    count = export_general_information(
        wb["General Information"], out_dir / "general_information.csv"
    )
    print(f"general_information.csv: {count} fields")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(Path(sys.argv[1]), Path(sys.argv[2]))
